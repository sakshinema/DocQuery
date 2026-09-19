from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import BinaryIO, Optional

from docx import Document as DocxDocument
from pypdf import PdfReader

from .normalize import ParsedPage, clean_text, remove_repeated_headers_footers

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv"}


class UnsupportedDocument(Exception):
    pass


class DocumentParseError(Exception):
    pass


def parse_document(filename: str, stream: BinaryIO) -> tuple[list[ParsedPage], dict]:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocument(f"Unsupported file type: {ext or 'unknown'}")
    try:
        data = stream.read()
        if not data:
            raise DocumentParseError("The file is empty.")
        if ext == ".pdf":
            pages, meta = _parse_pdf(data)
        elif ext == ".docx":
            pages, meta = _parse_docx(data)
        elif ext == ".csv":
            pages, meta = _parse_csv(data)
        else:
            pages, meta = _parse_text(data)
        pages = remove_repeated_headers_footers(pages)
        if not any(p.text.strip() for p in pages):
            raise DocumentParseError("No readable text was found. The document may be image-only or encrypted.")
        return pages, meta
    except UnsupportedDocument:
        raise
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(f"Could not parse document: {exc}") from exc


def _parse_pdf(data: bytes) -> tuple[list[ParsedPage], dict]:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise DocumentParseError("The PDF is password-protected and cannot be read.")
    pages: list[ParsedPage] = []
    for idx, page in enumerate(reader.pages, start=1):
        text = clean_text(page.extract_text() or "")
        pages.append(ParsedPage(idx, None, text))
    metadata = {}
    raw_meta = reader.metadata or {}
    if raw_meta:
        for key, value in raw_meta.items():
            if value:
                metadata[str(key).lstrip("/")] = str(value)
    return pages, {"page_count": len(reader.pages), "metadata": metadata}


def _parse_docx(data: bytes) -> tuple[list[ParsedPage], dict]:
    doc = DocxDocument(io.BytesIO(data))
    sections: list[str] = []
    current: str | None = None
    for paragraph in doc.paragraphs:
        text = clean_text(paragraph.text)
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        if "heading" in style:
            current = text
            sections.append(text)
        else:
            sections.append(f"{current or ''}@@{text}")
    # Convert the simple marker representation into section-aware page records.
    records: list[ParsedPage] = []
    current_section = None
    body_lines: list[str] = []
    for item in sections:
        if "@@" in item:
            prefix, body = item.split("@@", 1)
            current_section = prefix or current_section
            body_lines.append(body)
        else:
            if body_lines:
                records.append(ParsedPage(None, current_section, "\n".join(body_lines)))
                body_lines = []
            current_section = item
    if body_lines:
        records.append(ParsedPage(None, current_section, "\n".join(body_lines)))
    if not records:
        records = [ParsedPage(None, None, "\n".join(p.text for p in doc.paragraphs))]

    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append(" | ".join(clean_text(cell.text) for cell in row.cells))
        if rows:
            records.append(ParsedPage(None, "Table", "\n".join(rows)))
    props = doc.core_properties
    metadata = {"title": props.title or "", "author": props.author or ""}
    return records, {"page_count": None, "metadata": metadata}


def _parse_csv(data: bytes) -> tuple[list[ParsedPage], dict]:
    text = data.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return [], {"page_count": None, "metadata": {}}
    header = rows[0]
    lines = [" | ".join(cell.strip() for cell in header)]
    for row in rows[1:]:
        lines.append(" | ".join(cell.strip() for cell in row))
    return [ParsedPage(None, "CSV", "\n".join(lines))], {
        "page_count": None,
        "metadata": {"rows": len(rows) - 1, "columns": len(header)},
    }


def _parse_text(data: bytes) -> tuple[list[ParsedPage], dict]:
    text = data.decode("utf-8-sig", errors="replace")
    return [ParsedPage(None, None, clean_text(text))], {"page_count": None, "metadata": {}}
